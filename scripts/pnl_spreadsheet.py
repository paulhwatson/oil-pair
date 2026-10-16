#!/usr/bin/env python3
"""Each pair's profitability over time, as a Numbers spreadsheet.

Rebuilt from scratch on every run from IG's own records (oil_pair.ledger),
so it is always complete and never drifts from the account. Trades opened
before Mon 5 Oct (TRADES_FROM), under earlier versions of the strategy, are
left out.

  Summary          per pair: closed trades, won, P&L after costs banked,
                   the open position's P&L after costs now, and the total
  Trades           every closed pair trade: when, which way, IG's realised
                   P&L, financing and admin fees, and P&L after costs
  Over time        each pair's banked P&L after costs, cumulative, by day
  Daily snapshots  banked plus open P&L, recorded each run with --snapshot

The snapshots are the one part IG can't give back later - what an open
position was worth on a given evening - so each --snapshot run appends to
records/daily.csv and the sheet is drawn from that.

Usage:
    ./.venv/bin/python scripts/pnl_spreadsheet.py               # rebuild the spreadsheet
    ./.venv/bin/python scripts/pnl_spreadsheet.py --snapshot    # and record today's snapshot
    ./.venv/bin/python scripts/pnl_spreadsheet.py --dry-run     # print the summary only

Rebuilt weekdays at 18:05 by launchd, with a snapshot: see
scripts/install_pnl_spreadsheet.sh. Demo account only for now.
"""

from __future__ import annotations

import argparse
import logging
import shutil
import subprocess
import sys
import tomllib
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from oil_pair import costs, ledger  # noqa: E402
from oil_pair.ig_client import IGClient  # noqa: E402
from oil_pair.logging_setup import configure_logging  # noqa: E402
from oil_pair.settings import REPO_ROOT, load_credentials  # noqa: E402
from oil_pair.state_store import load_state  # noqa: E402
from oil_pair.strategy_logic import PairSide  # noqa: E402

UK = "Europe/London"
HISTORY_START = pd.Timestamp("2026-09-01", tz="UTC")  # the demo account's first deal
# Only trades opened from here count: Mon 5 Oct is when the strategy began
# running as designed (fits on one price an hour over 30 days, stops tuned on
# them, closes sized from IG). The earlier versions' trades are left out. The
# history is still read from HISTORY_START, as deals held then can share a
# later overnight charge with these.
TRADES_FROM = pd.Timestamp("2026-10-05", tz=UK)
DEFAULT_OUT = Path.home() / "Documents" / "oil-pair P&L.numbers"
SNAPSHOTS = REPO_ROOT / "records" / "daily.csv"
SNAPSHOT_COLUMNS = ["date", "pair", "account", "banked", "open", "total"]

log = logging.getLogger("pnl_spreadsheet")


def load_pairs() -> list[ledger.PairMarkets]:
    pairs = []
    for path in sorted((REPO_ROOT / "config").glob("*/pair_config.toml")):
        cfg = tomllib.loads(path.read_text())
        epics = {cfg[k]["epic"]: cfg[k]["name"] for k in ("instrument_a", "instrument_b")}
        pairs.append(ledger.PairMarkets(path.parent.name, epics))
    return pairs


def log_events(pairs: list[ledger.PairMarkets]) -> list[ledger.LogEvent]:
    events = []
    for pair in pairs:
        for path in sorted((REPO_ROOT / "logs" / pair.pair).glob("oil_pair.log*")):
            events += ledger.read_log(path, pair)
    return events


def closed_trades(pairs: list[ledger.PairMarkets]) -> list[ledger.Trade]:
    client = IGClient(load_credentials(live=False))
    client.login()
    transactions = costs.fetch_transactions(client, HISTORY_START)
    open_now = costs.holdings(client.fetch_open_positions())
    return ledger.build_trades(transactions, pairs, "demo", log_events(pairs), open_now)


def open_positions() -> dict[str, object]:
    """pair -> its PairStatus from the status report (P&L after costs if closed now)."""
    import pair_status  # scripts/pair_status.py
    statuses, _ = pair_status.gather(dry_run=True)
    return {s.name: s for s in statuses}


def summary_rows(pairs, trades, open_now, now) -> list[list]:
    rows = []
    for pair in pairs:
        mine = [t for t in trades if t.pair == pair.pair]
        paired = [t for t in mine if len(t.markets) == 2]  # lone legs count in banked, not as trades
        status = open_now.get(pair.pair)
        state = load_state(REPO_ROOT / "state" / pair.pair / "run_state.json")
        account = (status.account if status else (state.account if state and state.account else "demo")).upper()
        if status:
            since = status.opened.tz_convert(UK).strftime("%d %b")
            position = f"{status.side} since {since}"
        else:
            position = "flat" if state is None or state.side is PairSide.FLAT else state.side.value
        banked = sum(t.net for t in mine)
        open_pnl = status.net if status else 0.0
        won = sum(t.net > 0 for t in paired)
        rows.append([
            pair.pair, account, position, len(paired), won, (won / len(paired)) if paired else None,
            banked, sum(t.financing for t in mine), (sum(t.net for t in paired) / len(paired)) if paired else None,
            max((t.net for t in paired), default=None), min((t.net for t in paired), default=None),
            open_pnl, banked + open_pnl,
            min(t.opened for t in mine).tz_convert(UK).tz_localize(None).to_pydatetime() if mine else None,
        ])
    total = ["All pairs", "", "", sum(r[3] for r in rows), sum(r[4] for r in rows), None,
             sum(r[6] for r in rows), sum(r[7] for r in rows), None, None, None,
             sum(r[11] for r in rows), sum(r[12] for r in rows), None]
    if total[3]:
        total[5] = total[4] / total[3]
    return rows + [total]


def trade_rows(trades: list[ledger.Trade]) -> list[list]:
    running: dict[str, float] = {}
    rows = []
    for t in sorted(trades, key=lambda t: t.closed):
        running[t.pair] = running.get(t.pair, 0.0) + t.net
        rows.append([
            t.pair, t.account.upper(), t.side, _local(t.opened), _local(t.closed),
            (t.closed - t.opened).total_seconds() / 86400, t.exit, " / ".join(t.markets),
            t.realised, t.financing, t.net, running[t.pair], t.note,
        ])
    return rows[::-1]  # newest first


def over_time_rows(pairs, trades, now) -> tuple[list[str], list[list]]:
    names = [p.pair for p in pairs]
    first = TRADES_FROM.tz_localize(None)
    days = pd.date_range(first, now.tz_convert(UK).normalize().tz_localize(None), freq="D")
    banked = ledger.banked_by_day(trades, names, days)
    banked["All pairs"] = banked.sum(axis=1)
    return ["Date", *banked.columns], [[d.to_pydatetime(), *row] for d, row in zip(banked.index, banked.to_numpy().tolist())]


def record_snapshot(summary: list[list], now: pd.Timestamp) -> pd.DataFrame:
    day = now.tz_convert(UK).strftime("%Y-%m-%d")
    fresh = pd.DataFrame([[day, r[0], r[1], round(r[6], 2), round(r[11], 2), round(r[12], 2)] for r in summary[:-1]],
                         columns=SNAPSHOT_COLUMNS)
    old = pd.read_csv(SNAPSHOTS) if SNAPSHOTS.exists() else pd.DataFrame(columns=SNAPSHOT_COLUMNS)
    out = pd.concat([old[old.date != day], fresh], ignore_index=True).sort_values(["date", "pair"])
    SNAPSHOTS.parent.mkdir(parents=True, exist_ok=True)
    out.to_csv(SNAPSHOTS, index=False)
    return out


def _local(ts: pd.Timestamp):
    return ts.tz_convert(UK).tz_localize(None).to_pydatetime()


# --- writing the Numbers file ---
#
# The workbook is built with openpyxl and Numbers itself saves it as a
# .numbers file, so the file is one Numbers wrote and opens cleanly. (Writing
# .numbers directly with numbers-parser 4.22 gave every table the size of the
# last one created, so the Summary showed one column.)

MONEY, PERCENT, DATETIME, DATE, DAYS = "money", "percent", "datetime", "date", "days"
FORMATS = {
    # Numbers keeps this one as £ with pence and red losses; £#,##0.00 loses
    # both, and a [Red]-"£" negative section shows a doubled minus
    MONEY: '"£"#,##0.00;[Red]\\-"£"#,##0.00',
    PERCENT: "0%", DATETIME: "d mmm yyyy hh:mm", DATE: "ddd d mmm yyyy", DAYS: "0.0",
}
WORKBOOK = REPO_ROOT / "records" / "pnl.xlsx"


def write_sheet(ws, headers: list[str], rows: list[list], kinds: dict[int, str], widths: dict[int, int],
                bold_last_row: bool = False) -> None:
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
    for row in rows:
        ws.append([None if v == "" else round(v, 4) if isinstance(v, float) else v for v in row])
    for col, kind in kinds.items():
        for r in range(2, len(rows) + 2):
            ws.cell(r, col + 1).number_format = FORMATS[kind]
    if bold_last_row and rows:
        for cell in ws[len(rows) + 1]:
            cell.font = Font(bold=True)
    for col, header in enumerate(headers):
        ws.column_dimensions[get_column_letter(col + 1)].width = widths.get(col, max(11, len(header) + 2))
    ws.freeze_panes = "A2"


def build_workbook(summary, trades, over_time, snapshots, now: pd.Timestamp, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.utils import get_column_letter

    s_headers = ["Pair", "Account", "Now", "Closed trades", "Won", "Win rate", "Banked after costs",
                 "Financing in that", "Average per trade", "Best trade", "Worst trade", "Open P&L after costs",
                 "Total", "First trade"]
    t_headers = ["Pair", "Account", "Side", "Opened (UK)", "Closed (UK)", "Days held", "Exit", "Markets",
                 "Realised (IG)", "Financing", "P&L after costs", "Pair running total", "Note"]
    o_headers, o_rows = over_time
    d_headers = ["Date", "Pair", "Account", "Banked to date", "Open P&L after costs", "Total"]
    d_rows = [[pd.Timestamp(r.date).to_pydatetime(), r.pair, r.account, r.banked, r.open, r.total]
              for r in snapshots.itertuples()]
    notes = [
        [f"Updated {now.tz_convert(UK):%a %d %b %Y %H:%M} UK from IG's demo account history, rebuilt in full each time."],
        [f"Trades opened from {TRADES_FROM:%a %d %b %Y} only - when the strategy began running as designed "
         "(fits on one price an hour over 30 days). The earlier versions' trades on this account are left out."],
        ["Banked after costs = IG's realised P&L on each closed trade (the bid/offer is already in it) "
         "plus its share of IG's overnight financing and admin fees."],
        ["Open P&L after costs = what an open position would bank if closed at the last price: "
         "the move at mid, less the bid/offer, plus financing so far."],
        ["Closed trades, win rate, average, best and worst count pair trades only. A lone leg - an entry "
         "IG half-refused and the app unwound, or a leg left unpaired - is listed under Trades and counts "
         "in Banked."],
        ["Entries before Sat 10 Oct were sized before the central £165 risk setting, so smaller."],
        ["Daily snapshots start the day this sheet was set up; each weekday's 18:05 run adds one."],
        ["Edits here are overwritten at the next rebuild - copy a sheet out first to keep your own changes."],
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    write_sheet(ws, s_headers, summary, {5: PERCENT, 6: MONEY, 7: MONEY, 8: MONEY, 9: MONEY, 10: MONEY,
                                         11: MONEY, 12: MONEY, 13: DATE}, {0: 18, 2: 20, 13: 16}, bold_last_row=True)
    write_sheet(wb.create_sheet("Trades"), t_headers, trades,
                {3: DATETIME, 4: DATETIME, 5: DAYS, 8: MONEY, 9: MONEY, 10: MONEY, 11: MONEY},
                {0: 17, 3: 17, 4: 17, 7: 34, 12: 52})
    ws = wb.create_sheet("Over time")
    write_sheet(ws, o_headers, o_rows, {0: DATE, **{c: MONEY for c in range(1, len(o_headers))}},
                {0: 16, **{c: 15 for c in range(1, len(o_headers))}})
    if o_rows:
        chart = LineChart()
        chart.title = "Banked after costs, cumulative"
        chart.height, chart.width = 11, 26
        chart.x_axis.delete = chart.y_axis.delete = False  # openpyxl hides both axes unless told
        chart.x_axis.number_format, chart.y_axis.number_format = "d mmm", '"£"#,##0'
        chart.x_axis.tickLblSkip = 7
        last = len(o_rows) + 1
        # only pairs that have closed a trade - the rest would be flat lines at zero - and the total
        traded = [c for c in range(1, len(o_headers) - 1) if any(row[c] for row in o_rows)]
        for c in [*traded, len(o_headers) - 1]:
            chart.add_data(Reference(ws, min_col=c + 1, min_row=1, max_row=last), titles_from_data=True)
        total = chart.series[-1]
        total.graphicalProperties.line.solidFill = "000000"
        total.graphicalProperties.line.width = 28000  # EMU, about 2.2pt
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
        ws.add_chart(chart, f"{get_column_letter(len(o_headers) + 2)}2")
    write_sheet(wb.create_sheet("Daily snapshots"), d_headers, d_rows, {0: DATE, 3: MONEY, 4: MONEY, 5: MONEY},
                {0: 16, 1: 18, 4: 20})
    write_sheet(wb.create_sheet("Notes"), ["About this spreadsheet"], notes, {}, {0: 150})
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _applescript_string(path: Path) -> str:
    return '"' + str(path).replace("\\", "\\\\").replace('"', '\\"') + '"'


def save_as_numbers(workbook: Path, out: Path) -> None:
    """Have Numbers open the workbook and save it as `out`, then close it -
    and quit Numbers again if it wasn't already running. Saved beside the
    workbook first and moved into place, so a half-written file never
    replaces the last good one."""
    staged = workbook.with_suffix(".numbers")
    if staged.is_dir():
        shutil.rmtree(staged)
    staged.unlink(missing_ok=True)
    script = f"""
    set wasRunning to application "Numbers" is running
    tell application "Numbers"
        set d to open POSIX file {_applescript_string(workbook)}
        save d in POSIX file {_applescript_string(staged)}
        close d saving no
        if not wasRunning then quit
    end tell
    """
    subprocess.run(["osascript", "-e", script], check=True, capture_output=True, text=True, timeout=180)
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.is_dir():
        shutil.rmtree(out)
    shutil.move(str(staged), str(out))


def main() -> None:
    parser = argparse.ArgumentParser(description="Each pair's profitability over time, as a Numbers spreadsheet.")
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT, help=f"where to write it (default: {DEFAULT_OUT})")
    parser.add_argument("--snapshot", action="store_true", help="also record today's banked and open P&L")
    parser.add_argument("--dry-run", action="store_true", help="print the summary instead of writing")
    args = parser.parse_args()
    from dotenv import load_dotenv
    load_dotenv(REPO_ROOT / ".env")
    if not args.dry_run:
        configure_logging(log_dir=REPO_ROOT / "logs" / "pnl_spreadsheet")

    now = pd.Timestamp.now(tz="UTC")
    pairs = load_pairs()
    trades = [t for t in closed_trades(pairs) if t.opened >= TRADES_FROM]
    summary = summary_rows(pairs, trades, open_positions(), now)
    if args.dry_run:
        for row in summary:
            print(f"{row[0]:<17} {row[2]:<20} trades {row[3]:>3}  banked £{row[6]:>9,.2f}  open £{row[11]:>8,.2f}  "
                  f"total £{row[12]:>9,.2f}")
        return
    snapshots = record_snapshot(summary, now) if args.snapshot else (
        pd.read_csv(SNAPSHOTS) if SNAPSHOTS.exists() else pd.DataFrame(columns=SNAPSHOT_COLUMNS))
    build_workbook(summary, trade_rows(trades), over_time_rows(pairs, trades, now), snapshots, now, WORKBOOK)
    try:
        save_as_numbers(WORKBOOK, args.out)
    except (subprocess.SubprocessError, OSError) as exc:
        # e.g. macOS hasn't allowed this process to control Numbers yet: keep the workbook, which Numbers opens too
        fallback = args.out.with_suffix(".xlsx")
        shutil.copy(WORKBOOK, fallback)
        detail = getattr(exc, "stderr", "") or exc
        log.error("Numbers couldn't save %s (%s) - wrote %s instead", args.out, str(detail).strip(), fallback)
        print(f"Numbers couldn't save it - wrote {fallback} instead")
        sys.exit(1)
    log.info("wrote %s: %d closed trades, %d pairs%s", args.out, len(trades), len(pairs),
             ", with today's snapshot" if args.snapshot else "")
    print(f"wrote {args.out}")


if __name__ == "__main__":
    main()
